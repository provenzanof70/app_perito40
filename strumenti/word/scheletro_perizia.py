#!/usr/bin/env python3
"""Scheletro di perizia giurata e analisi tecnica in Word, generato dal fascicolo della pratica.

Ogni campo è compilato dal fascicolo (pratica_40.json e beni/Bnn.json) oppure segnato in chiaro:
  [PLACEHOLDER: …]  dato mancante, con la fonte attesa
  [NOTA: …]         istruzione per il perito
  [ev: file, p.X]   rimando all'evidenza
Le conclusioni e le caselle dei requisiti non si spuntano mai da sole: le decide il perito.

Uso:
  python strumenti/word/scheletro_perizia.py <cartella pratica> [--bene B01] [--perito perito.json] [--uscita cartella] [--data AAAA-MM-GG]
Senza --uscita scrive in <cartella pratica>/perizia/. Senza --perito i dati del perito restano segnaposto.
Il file del perito (privato, mai nel repository) ha i campi: nome, titolo, albo, ordine, natoA, natoIl, codiceFiscale,
residenza, tribunale, luogo, intestazione.
Richiede Python 3 con python-docx e Node.js (per leggere la guida ai requisiti dell'app).
"""

import argparse
import datetime as dt
import json
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

RADICE = Path(__file__).resolve().parents[2]
GRIGIO = RGBColor(0x59, 0x59, 0x59)
ROSSO = RGBColor(0xB0, 0x24, 0x18)
EVENTI = {
    'consegna': ('ddt_consegna',), 'collaudo': ('collaudo_sat',), 'funzione': ('entrata_in_funzione',),
    'interconnessione': ('interconnessione_dichiarata', 'prima_prova_scambio'), 'verifica': ('verifica_perito',),
}


# ---------------------------------------------------------------- dati

def leggi_json(p):
    return json.loads(Path(p).read_text(encoding='utf-8'))


def guida_requisiti():
    codice = "import('./app/js/guida.js').then(m => console.log(JSON.stringify(m.GUIDA_REQUISITI)))"
    r = subprocess.run(['node', '--input-type=module', '-e', codice], cwd=RADICE, capture_output=True, text=True, encoding='utf-8', check=True)
    return json.loads(r.stdout)


def ph(testo):
    return f'[PLACEHOLDER: {testo}]'


def rif_ev(fonti):
    return ' '.join(f"[ev: {f.get('file') or f.get('idEvidenza')}{', ' + f['posizione'] if f.get('posizione') else ''}]" for f in fonti or [])


def bene_di(codice):
    return codice.split('.')[0]


def data_it(iso):
    if not iso:
        return None
    parti = iso[:10].split('-')
    return '/'.join(reversed(parti)) if len(parti) == 3 else iso


def euro(v):
    return ph('importo') if v is None else f"€ {v:,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.')


def voce(scheda, sezione, campo, cosa):
    """Valore agli atti con fonte ed esito in campo; segnaposto se manca."""
    v = ((scheda.get('datiAgliAtti') or {}).get(sezione) or {}).get(campo)
    if not v or v.get('nonApplicabile') and v.get('valore') is None:
        return 'non applicabile' if v and v.get('nonApplicabile') else ph(f'{cosa} (documenti del bene)')
    valore = v.get('valore')
    vc = v.get('verificaCampo') or {}
    if vc.get('esito') == 'corretto' and vc.get('valoreRilevato'):
        valore = f"{vc['valoreRilevato']} (rilevato in campo; agli atti: {v.get('valore')})"
    if valore is None:
        return ph(f'{cosa}: non indicato nei documenti')
    testo = f'{valore} {rif_ev(v.get("fonti"))}'.strip()
    if vc.get('esito') == 'difforme':
        testo += f" [NOTA: DIFFORMITÀ in campo: {vc.get('valoreRilevato', '')} — la bozza non è giurabile finché non è risolta]"
    elif v.get('stato') != 'documentato':
        testo += f" [NOTA: dato {v.get('stato', 'da verificare').replace('_', ' ')}]"
    return testo


class Pratica:
    def __init__(self, cartella):
        self.cartella = Path(cartella)
        self.p = leggi_json(self.cartella / 'pratica_40.json')
        self.schede = {b['codice']: leggi_json(self.cartella / b['scheda']) for b in self.p['beni']}

    def eventi(self, codice, chiave):
        tipi = EVENTI[chiave]
        righe = [t for t in self.p['timeline'] if t['evento'] in tipi and codice in [bene_di(b) for b in t['beni']]]
        return sorted(righe, key=lambda t: t['data']['valore'])

    def evento(self, codice, chiave, cosa):
        righe = self.eventi(codice, chiave)
        if not righe:
            return ph(f'data di {cosa} (timeline)')
        t = righe[0]
        testo = f"{data_it(t['data']['valore'])} ({t['stato']}) {rif_ev(t.get('prove'))}".strip()
        return testo

    def evidenze(self, codice, requisito):
        return [e for e in self.p['evidenze'] if e.get('stato', 'valida') == 'valida' and codice in [bene_di(b) for b in e['beni']] and requisito in e['requisiti']]

    def riscontri(self, codice, requisito):
        return [r for r in self.p['riscontriDiretti'] if r['bene'] == codice and r['requisito'] == requisito]

    def dichiarazioni(self, codice, requisito):
        return [d for d in self.p['dichiarazioniTerzi'] if codice in d['beni'] and d['requisito'] == requisito]

    def azioni(self, codice, requisito=None):
        return [a for a in self.p['azioni'] if a['stato'] == 'aperta' and codice in [bene_di(b) for b in a['beni']] and (requisito is None or a.get('requisito') == requisito)]

    def ultima_verifica(self):
        eseguiti = [s for s in self.p['sopralluoghi'] if s['stato'] == 'eseguito']
        return data_it(max(s['data'] for s in eseguiti)) if eseguiti else None


# ---------------------------------------------------------------- Word

class Doc:
    def __init__(self, titolo_piede, perito):
        self.d = Document()
        st = self.d.styles['Normal']
        st.font.name = 'Arial'
        st.font.size = Pt(10.5)
        st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
        for nome, dim in (('Heading 1', 15), ('Heading 2', 12.5), ('Heading 3', 11), ('Title', 17)):
            s = self.d.styles[nome]
            s.font.name = 'Arial'
            s.font.size = Pt(dim)
            s.font.color.rgb = RGBColor(0x1F, 0x1F, 0x1F)
            f = s.element.rPr.rFonts
            f.set(qn('w:eastAsia'), 'Arial')
            # i titoli di Word usano il carattere del tema: si toglie perché valga Arial
            for attr in ('w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme'):
                f.attrib.pop(qn(attr), None)
        sez = self.d.sections[0]
        sez.orientation = WD_ORIENT.PORTRAIT
        sez.page_width, sez.page_height = Cm(21), Cm(29.7)
        sez.left_margin = sez.right_margin = sez.bottom_margin = Cm(2.0)
        sez.top_margin = Cm(2.3)
        testa = sez.header.paragraphs[0]
        testa.text = perito.get('intestazione') or ph('intestazione del perito (file del perito)')
        testa.runs[0].font.size = Pt(8.5)
        piede = sez.footer.paragraphs[0]
        r = piede.add_run(titolo_piede + '    Pag. ')
        r.font.size = Pt(8)
        self._campo(piede, 'PAGE')
        piede.add_run(' di ').font.size = Pt(8)
        self._campo(piede, 'NUMPAGES')

    @staticmethod
    def _campo(par, nome):
        run = par.add_run()
        run.font.size = Pt(8)
        for tipo, testo in (('begin', None), (None, nome), ('end', None)):
            if tipo:
                el = OxmlElement('w:fldChar')
                el.set(qn('w:fldCharType'), tipo)
            else:
                el = OxmlElement('w:instrText')
                el.set(qn('xml:space'), 'preserve')
                el.text = testo
            run._r.append(el)

    def p(self, testo='', grassetto=False, centro=False, corsivo=False, dim=None):
        par = self.d.add_paragraph()
        self._testo(par, testo, grassetto, corsivo, dim)
        if centro:
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        return par

    def _testo(self, par, testo, grassetto=False, corsivo=False, dim=None):
        # i segnaposto in rosso, le note in grigio corsivo: si vedono subito nella bozza
        import re
        for pezzo in re.split(r'(\[PLACEHOLDER:[^\]]*\]|\[NOTA:[^\]]*\])', testo):
            if not pezzo:
                continue
            r = par.add_run(pezzo)
            r.bold, r.italic = grassetto, corsivo
            if dim:
                r.font.size = Pt(dim)
            if pezzo.startswith('[PLACEHOLDER'):
                r.font.color.rgb = ROSSO
            elif pezzo.startswith('[NOTA'):
                r.font.color.rgb, r.italic = GRIGIO, True

    def h(self, testo, livello=1):
        return self.d.add_heading(testo, livello)

    def elenco(self, righe):
        for t in righe:
            self._testo(self.d.add_paragraph(style='List Bullet'), t)

    def casella(self, testo, spuntata=False):
        self.p(f"{'☒' if spuntata else '☐'} {testo}")

    def tabella(self, righe, intestazione=None, larghezze=None):
        t = self.d.add_table(rows=0, cols=len(righe[0] if righe else intestazione))
        t.style = 'Table Grid'
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, riga in enumerate(([intestazione] if intestazione else []) + righe):
            celle = t.add_row().cells
            for j, valore in enumerate(riga):
                celle[j].paragraphs[0].text = ''
                self._testo(celle[j].paragraphs[0], str(valore), grassetto=bool(intestazione) and i == 0, dim=9.5)
                if larghezze:
                    celle[j].width = Cm(larghezze[j])
        self.d.add_paragraph()
        return t

    def riquadro(self, testo):
        t = self.d.add_table(rows=1, cols=1)
        t.style = 'Table Grid'
        self._testo(t.rows[0].cells[0].paragraphs[0], testo, grassetto=True)
        self.d.add_paragraph()

    def salto(self):
        self.d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def salva(self, percorso):
        self.d.save(percorso)


# ---------------------------------------------------------------- contenuti comuni

def dati_comuni(pr, sc, perito, oggi):
    c = pr.p['cliente']
    el = sc.get('elaborato') or {}
    rev = el.get('revisione', 0)
    stato = el.get('stato', 'bozza')
    return {
        'cliente': c['ragioneSociale'], 'bene': sc['denominazione'], 'rev': rev, 'stato': stato,
        'bozza': stato in ('bozza', 'bozza_sospesa'), 'data': data_it(el.get('data')) or oggi,
        'sedeLegale': c.get('sedeLegale') or ph('sede legale (visura camerale)'),
        'piva': c.get('partitaIva') or ph('partita IVA (visura camerale)'),
        'legale': (c.get('legaleRappresentante') or {}).get('nome') or ph('legale rappresentante (visura camerale)'),
        'carica': (c.get('legaleRappresentante') or {}).get('carica') or ph('carica del legale rappresentante'),
        'perito': f"{perito.get('titolo') or ''} {perito['nome']}".strip() if perito.get('nome') else ph('nome del perito (file del perito)'),
        'tribunale': perito.get('tribunale') or ph('tribunale (file del perito)'),
        'luogo': perito.get('luogo') or ph('luogo'),
        'giuramento': data_it(((el.get('giuramento') or {}).get('data'))) or ph('data del giuramento'),
    }


def riquadro_bozza(doc, k):
    if k['bozza']:
        sospesa = ' SOSPESA' if k['stato'] == 'bozza_sospesa' else ''
        doc.riquadro(f"BOZZA{sospesa} DI LAVORO – non giurabile allo stato. Revisione {k['rev']} del {k['data']}, generata dal fascicolo: "
                     'segnaposto e note vanno risolti e tolti prima della firma.')


def copertina(doc, pr, sc, k, titolo, sottotitolo, regime):
    c = pr.p['cliente']
    logo = c.get('logo')
    if logo and (pr.cartella / logo).exists():
        doc.d.add_picture(str(pr.cartella / logo), width=Cm(4))
    else:
        doc.p(ph("logo dell'impresa (cliente.logo nel fascicolo)"))
    doc.p(k['cliente'], grassetto=True, dim=13)
    doc.p(k['sedeLegale'])
    doc.p(f"P.IVA {k['piva']}")
    doc.p()
    doc.p(titolo, grassetto=True, centro=True, dim=17)
    doc.p(sottotitolo, centro=True)
    doc.p()
    ridotto = sc['classificazione']['profilo'] == 'ridotto'
    doc.p(('BENE ' if ridotto else 'BENE MATERIALE composto da: ') + sc['denominazione'], grassetto=True, centro=True)
    for comp in sc.get('componenti', []):
        doc.p(f"{comp['codice']} {comp['denominazione']}", centro=True, dim=9.5)
    doc.p()
    cl = sc['classificazione']
    doc.p(f"Allegato {cl['allegato']}{' – Gruppo ' + str(cl['gruppo']) if cl.get('gruppo') else ''} annesso alla Legge 11 dicembre 2016, n. 232", centro=True)
    doc.p('Beni funzionali alla trasformazione tecnologica e digitale delle imprese secondo il modello «Industria 4.0»', centro=True, corsivo=True)
    doc.p()
    doc.tabella([['Perito giurato', k['perito']], ['Data giuramento', k['giuramento']], ['Ente asseveratore', k['tribunale']]], larghezze=[5, 11])
    doc.salto()


# ---------------------------------------------------------------- perizia

def perizia(pr, codice, perito, regime, oggi):
    sc = pr.schede[codice]
    k = dati_comuni(pr, sc, perito, oggi)
    cl = sc['classificazione']
    piede = f"Perizia 4.0 – {k['cliente']} – Bene {k['bene']} – {'BOZZA ' if k['bozza'] else ''}rev. {k['rev']} del {k['data']}"
    doc = Doc(piede, perito)
    copertina(doc, pr, sc, k, 'PERIZIA TECNICA GIURATA', f"ai sensi dell'{regime['riferimentoAnalisi']}", regime)
    riquadro_bozza(doc, k)
    doc.p('PERIZIA GIURATA', grassetto=True, centro=True, dim=15)
    nato = f"nato a {perito['natoA']} il {data_it(perito['natoIl'])}" if perito.get('natoA') else ph('luogo e data di nascita del perito')
    doc.p(f"Il sottoscritto {k['perito']}, {nato}, C.F. {perito.get('codiceFiscale') or ph('codice fiscale')}, "
          f"residente in {perito.get('residenza') or ph('residenza')}, iscritto all'Albo {perito.get('albo') or ph('iscrizione all’albo')}, "
          f"incaricato dalla società {k['cliente']}, con sede legale in {k['sedeLegale']}, P.IVA {k['piva']},")
    doc.p('VISTI', grassetto=True, centro=True)
    doc.elenco(regime['visti'])
    doc.p('[NOTA: elenco dei VISTI da verificare per regime e annualità della pratica]')
    doc.p('PREMESSO', grassetto=True, centro=True)
    doc.p(f"che la società {k['cliente']}, in persona del legale rappresentante {k['legale']} ({k['carica']}), dichiara:")
    valore = (sc.get('costi') or {}).get('valoreComplessivo')
    doc.elenco(regime['premesso'] + [f"che il valore complessivo dell'investimento è di {euro(valore)}"])
    doc.p('VISTO il bene', grassetto=True, centro=True)
    forn = sc.get('fornitore')
    doc.tabella([
        ['PRODUTTORE', voce(sc, 'identita', 'costruttore', 'costruttore')],
        ['MODELLO', voce(sc, 'identita', 'modello', 'modello')],
        ['MATRICOLA', voce(sc, 'identita', 'matricola', 'matricola')],
        ['DESCRIZIONE', sc['funzione'] + ('; componenti: ' + ', '.join(c['denominazione'] for c in sc['componenti']) if sc.get('componenti') else '')],
        ['ACQUISTATO DA', f"{k['cliente']} – {k['sedeLegale']}"],
        ['FORNITORE', f"{forn['nome']} ({forn['ruolo']}) {rif_ev(forn.get('fonti'))}".strip() if forn else ph('fornitore o integratore (ordine, fattura)')],
        ['SITO DI MESSA IN FUNZIONE', sc.get('unitaLocale') or pr.p['cliente']['sedeIntervento'] + (f" – {sc['ubicazione']}" if sc.get('ubicazione') else '')],
    ], larghezze=[5, 11])
    verifica = pr.ultima_verifica() or ph('data dell’ultimo sopralluogo eseguito')
    doc.p(f'così come risultante dalle attività di verifica ultimate in data {verifica},')
    doc.p('DICHIARA', grassetto=True, centro=True)
    doc.p('[NOTA: le caselle si spuntano solo a esito positivo della matrice RO/RU dell’Analisi tecnica]')
    voce_num = cl.get('codiceVoce') or ph('voce numerata (es. A11_01)')
    doc.casella(f"che il bene rientra nell'Allegato {cl['allegato']}{', Gruppo ' + str(cl['gruppo']) if cl.get('gruppo') else ''}, voce {voce_num}: {cl['voce']} {rif_ev(cl.get('fonti'))}")
    if cl['profilo'] == 'completo':
        doc.p('e che soddisfa i seguenti requisiti obbligatori:')
        for r, t in regime['obbligatori'].items():
            doc.casella(f'{r.replace("RO", "O")}. {t}')
        doc.p('e almeno due dei seguenti requisiti ulteriori:')
        for r, t in regime['ulteriori'].items():
            doc.casella(f'{r.replace("RU", "U")}. {t}')
    else:
        doc.p('SODDISFA il requisito obbligatorio di INTERCONNESSIONE   ☐ SI   ☐ NO')
        doc.p(regime['interconnessioneRidotto'] + '.')
    doc.p(f"MESSO IN FUNZIONE: {pr.evento(codice, 'funzione', 'entrata in funzione')}")
    doc.p(f"INTERCONNESSO: {pr.evento(codice, 'interconnessione', 'interconnessione')} [NOTA: mai anteriore alla prima prova datata dello scambio]")
    doc.p(f"L'interconnessione è stata verificata in data {verifica}.")
    doc.p('E DICHIARA INOLTRE', grassetto=True, centro=True)
    doc.elenco(regime['dichiaraInoltre'])
    doc.p(f"{k['luogo']}, {k['data']}")
    doc.p(f"Timbro e firma del tecnico incaricato – {k['perito']}")
    doc.salto()
    doc.h('ACCORDO DI RISERVATEZZA (NDA)', 1)
    doc.p(f"tra {k['cliente']} (Committente), in persona del legale rappresentante {k['legale']}, e {k['perito']} (Perito).")
    doc.elenco(regime['nda'])
    doc.tabella([['Per la Committente', 'Per il Perito'], [k['legale'], k['perito']]], larghezze=[8, 8])
    doc.salto()
    doc.h('VERBALE DI ASSEVERAZIONE', 1)
    g = (sc.get('elaborato') or {}).get('giuramento') or {}
    doc.p(f"R.G. V.G. N° {g.get('rgvg') or ph('dopo il giuramento')}   CRON. N° {g.get('cronologico') or ph('dopo il giuramento')}")
    doc.p(f"{k['tribunale']} – Verbale di asseverazione di perizia stragiudiziale")
    doc.p(f"In data {k['giuramento']} è comparso {k['perito']}, iscritto all'Albo {perito.get('albo') or ph('iscrizione')}, "
          f"che chiede di asseverare con giuramento l'elaborato del {k['data']} redatto nell'interesse di {k['cliente']}. "
          "Ammonito ai sensi dell'art. 193 c.p.c. e dell'art. 483 c.p., presta giuramento ripetendo:")
    doc.p(f"«{regime['formulaGiuramento']}».", corsivo=True, centro=True)
    doc.tabella([['Il perito', 'Il funzionario'], [k['perito'], '']], larghezze=[8, 8])
    return doc, k


# ---------------------------------------------------------------- analisi tecnica

def sezione_requisito(doc, pr, codice, req, titolo, guida):
    doc.h(f'{req} {titolo}', 3)
    ev = pr.evidenze(codice, req)
    rd = pr.riscontri(codice, req)
    dt_ = pr.dichiarazioni(codice, req)
    if ev:
        doc.p('Evidenze:', grassetto=True)
        doc.elenco([f"{e['id']} – {e.get('descrizione') or e['tipo']} [ev: {Path(e['file']).name}]" for e in ev])
    if rd:
        doc.p('Riscontri diretti del perito:', grassetto=True)
        doc.elenco([f"{r['id']}: {r['frase']} [ev: {', '.join(r['evidenze'])}]" for r in rd])
    if dt_:
        doc.p('Dichiarazioni di terzi:', grassetto=True)
        doc.elenco([f"{d['id']} – {d['persona']} ({d['ruolo']}), {d['forma']}: «{d['contenuto']}» – riscontro oggettivo: {d['riscontroOggettivo']['esito']}" for d in dt_])
    if not (ev or rd or dt_):
        g = guida.get(req) or {}
        doc.p(ph('evidenze del requisito: nessuna nel fascicolo'))
        if g.get('cercare'):
            doc.p('[NOTA: da acquisire, ad esempio: ' + '; '.join(g['cercare'][:3]) + ']')
    lacune = pr.azioni(codice, req)
    if lacune:
        doc.p('Lacune:', grassetto=True)
        doc.elenco([f"{a['id']} ({a['priorita']}): {a['cosaServe']}" for a in lacune])
    doc.p('Formula conclusiva (allo stato): ' + ph('positiva / parziale / debole / negativa, decisa dal perito'))


def analisi(pr, codice, perito, regime, guida, oggi):
    sc = pr.schede[codice]
    k = dati_comuni(pr, sc, perito, oggi)
    cl = sc['classificazione']
    ridotto = cl['profilo'] == 'ridotto'
    piede = f"ANALISI TECNICA a corredo Perizia 4.0 – {k['cliente']} – Bene {k['bene']} – {'BOZZA ' if k['bozza'] else ''}rev. {k['rev']} del {k['data']}"
    doc = Doc(piede, perito)
    copertina(doc, pr, sc, k, 'ANALISI TECNICA', f"a corredo della perizia giurata ex {regime['riferimentoAnalisi']}", regime)
    riquadro_bozza(doc, k)

    doc.h('1 ANALISI TECNICA – REQUISITI', 1)
    if ridotto:
        doc.p('Per i beni dell’Allegato A, Gruppi II e III, e dell’Allegato B il requisito obbligatorio è l’interconnessione: ' + regime['interconnessioneRidotto'] + '.')
    else:
        doc.h('1.1.1 Requisiti obbligatori', 2)
        doc.p('Il bene deve possedere tutte le seguenti caratteristiche (Circolare 4/E, par. 11.1). È ammessa la perizia in due fasi: '
              'verifica dei requisiti tecnici e, a seguire, verifica dell’avvenuta interconnessione (Circolare 4/E, par. 6.3).')
        for r, t in regime['obbligatori'].items():
            doc.p(f'{r}. {t}', grassetto=True)
            doc.p((guida.get(r) or {}).get('norma', ''))
        doc.h('1.1.2 Requisiti ulteriori', 2)
        doc.p('Il bene deve possedere almeno due delle seguenti caratteristiche (Circolare 4/E, par. 11.2).')
        for r, t in regime['ulteriori'].items():
            doc.p(f'{r}. {t}', grassetto=True)
            doc.p((guida.get(r) or {}).get('norma', ''))

    doc.h('2 ANALISI TECNICA DEL BENE OGGETTO DI VALUTAZIONE', 1)
    doc.h('2.1 Descrizione', 2)
    doc.p(sc['funzione'] + '. ' + ph('funzione effettiva, targa letta, flusso funzionale dichiarato e osservato, figure con didascalia e [ev:]'))
    forn = sc.get('fornitore')
    comp = '; '.join(f"{c['codice']} {c['denominazione']}" for c in sc.get('componenti', [])) or 'nessun componente censito'
    doc.p('DATI IDENTIFICATIVI DEL BENE', grassetto=True)
    doc.tabella([
        ['Fornitore / integratore', f"{forn['nome']} ({forn['ruolo']})" if forn else ph('fornitore o integratore')],
        ['Costruttore', voce(sc, 'identita', 'costruttore', 'costruttore')],
        ['Modello e matricola', f"{voce(sc, 'identita', 'modello', 'modello')} – matricola {voce(sc, 'identita', 'matricola', 'matricola')}"],
        ['Tipologia', sc['funzione']],
        ['Componenti', comp],
        ['Anno di fabbricazione', voce(sc, 'identita', 'anno', 'anno')],
        ['Controllo (PLC/CNC)', f"{voce(sc, 'controllore', 'marca', 'marca del controllore')} {voce(sc, 'controllore', 'modello', 'modello del controllore')}"],
        ['Consegna', pr.evento(codice, 'consegna', 'consegna')],
        ['Collaudo', pr.evento(codice, 'collaudo', 'collaudo')],
        ['Entrata in funzione', pr.evento(codice, 'funzione', 'entrata in funzione')],
        ['Interconnessione', pr.evento(codice, 'interconnessione', 'interconnessione')],
    ], larghezze=[5, 11])

    doc.h('2.2 Classificazione', 2)
    voce_num = cl.get('codiceVoce') or ph('voce numerata')
    doc.p(f"Il bene rientra nell'Allegato {cl['allegato']}{', Gruppo ' + str(cl['gruppo']) if cl.get('gruppo') else ''}, voce {voce_num}: «{cl['voce']}», "
          f"in quanto {cl['motivazione']} {rif_ev(cl.get('fonti'))}")
    gate = (sc.get('datiAgliAtti') or {}).get('esitoGateC')
    if gate:
        doc.p(f"Perimetro e conformità dell'insieme (GATE C): esito {gate['esito']} – {gate['motivazione']} {rif_ev(gate.get('fonti'))}")
    doc.p(f"Grado di sostenibilità della classificazione: {cl['sostenibilita'].replace('_', ' ')}.")

    doc.h('2.3 Costo del bene e dei suoi componenti e accessori (così come risultante dalle fatture)', 2)
    costi = sc.get('costi') or {}
    doc.p(f"Valore complessivo dell'investimento: {euro(costi.get('valoreComplessivo'))}.")
    if costi.get('documenti'):
        doc.tabella([[f"{d['tipo'].replace('_', ' ')} n. {d['numero']}", data_it(d['data']), d['descrizione'] + ' ' + rif_ev([d['fonte']] if d.get('fonte') else []), euro(d['imponibile'])]
                     for d in costi['documenti']], intestazione=['Documento', 'Data', 'Descrizione', 'Imponibile'], larghezze=[4, 2.5, 7, 3])
    else:
        doc.p(ph('documenti di costo: tipo, numero, data, descrizione, imponibile (fatture, ordine, leasing)'))
    if costi.get('acconti'):
        tot = sum(a['importo'] for a in costi['acconti'])
        perc = f" pari al {tot / costi['valoreComplessivo'] * 100:.1f}% del valore".replace('.', ',') + f" (soglia {regime['sogliaAcconto'] * 100:.0f}%)" if costi.get('valoreComplessivo') else ''
        doc.p(f"Acconti: {', '.join(euro(a['importo']) + ' del ' + data_it(a['data']) for a in costi['acconti'])}{perc}. [NOTA: verificare la data limite dell'acconto per l'annualità]")
    dl = costi.get('dicituraLegge')
    doc.p(f"Dicitura di legge in fattura: {('presente – «' + dl.get('testo', '') + '» ' + rif_ev([dl['fonte']] if dl.get('fonte') else [])) if dl and dl['presente'] else ph('presente / assente')}")
    doc.p(regime['notaAccessori'], corsivo=True, dim=9)

    if ridotto:
        doc.h('2.4 Esame del requisito di interconnessione', 2)
        sezione_requisito(doc, pr, codice, 'RO2', 'Interconnessione', guida)
        doc.h('2.5 Riconducibilità alla voce', 2)
        sezione_requisito(doc, pr, codice, 'VOCE', 'Caratteristiche della voce', guida)
    else:
        doc.h('2.4 Esame dei requisiti obbligatori', 2)
        for r, t in regime['obbligatori'].items():
            sezione_requisito(doc, pr, codice, r, t, guida)
        doc.h('2.5 Esame dei requisiti ulteriori', 2)
        dichiarati = cl.get('ruDichiarati') or []
        for r, t in regime['ulteriori'].items():
            sezione_requisito(doc, pr, codice, r, t + ('' if r in dichiarati else ' [NOTA: non dichiarato dal cliente]'), guida)
        doc.h('2.6 Control check', 2)
        doc.p('[NOTA: le check list riportano lo stato alla data della bozza]')
        doc.p('Check List 1 – Voci dell’Allegato A, Gruppo I (Circ. 4/E, par. 11)', grassetto=True)
        for cod, t in regime['vociGruppo1'].items():
            doc.casella(f'{cod} {t}', cod == cl.get('codiceVoce'))
        doc.p('Check List 2 – Requisiti obbligatori', grassetto=True)
        doc.tabella([[r, t, ph('esito')] for r, t in regime['obbligatori'].items()], intestazione=['Req.', 'Caratteristica', 'Esito'], larghezze=[1.5, 11, 3.5])
        doc.p('Check List 3 – Requisiti ulteriori', grassetto=True)
        doc.tabella([[r, t, ph('esito')] for r, t in regime['ulteriori'].items()], intestazione=['Req.', 'Caratteristica', 'Esito'], larghezze=[1.5, 11, 3.5])

    doc.h('3 RAPPRESENTAZIONE DEI FLUSSI', 1)
    doc.p(ph('diagramma a blocchi: flusso dei materiali e flusso delle informazioni con colori distinti'))
    doc.h('4 CONCLUSIONI', 1)
    doc.p('[NOTA: da redigere solo a esito positivo della matrice RO/RU]')
    doc.p(f"Per quanto indicato sopra, con la presente Analisi tecnica si attesta che il bene {k['bene']} " + ph('conclusione'))
    doc.p(f"{k['luogo']}, {k['data']} – {k['perito']}")
    doc.h('5 ALLEGATI', 1)
    doc.elenco(regime['allegati'])

    doc.salto()
    doc.h('APPENDICE ISTRUTTORIA – documento interno, da rimuovere prima della firma', 1)
    reqs = ['RO2', 'VOCE'] if ridotto else list(regime['obbligatori']) + list(regime['ulteriori'])
    doc.p('A. Matrice dei requisiti', grassetto=True)
    righe = []
    for r in reqs:
        ev = pr.evidenze(codice, r)
        gap = '; '.join(a['cosaServe'] for a in pr.azioni(codice, r)) or ''
        righe.append([r, ev[0]['id'] if ev else ph('primaria'), ', '.join(e['id'] for e in ev[1:3]) or '', ph('0–5'), gap, ph('rischio'), ph('conclusione')])
    doc.tabella(righe, intestazione=['Req.', 'Evidenza primaria', 'Corroborante', 'Forza', 'Gap residuo', 'Rischio', 'Conclusione'])
    doc.p('B. Timeline di coerenza', grassetto=True)
    tl = sorted([t for t in pr.p['timeline'] if codice in [bene_di(b) for b in t['beni']]], key=lambda t: t['data']['valore'])
    if tl:
        doc.tabella([[data_it(t['data']['valore']) + f" ({t['data']['precisione']})", t['evento'].replace('_', ' '), rif_ev(t.get('prove')) or ph('prova'), t.get('coerenza', ''), t.get('nota', '')]
                     for t in tl], intestazione=['Data', 'Evento', 'Documento di prova', 'Coerenza', 'Nota'])
    else:
        doc.p(ph('timeline del bene: nessuna data nel fascicolo'))
    doc.p('C. Integrazioni richieste', grassetto=True)
    az = pr.azioni(codice)
    if az:
        doc.tabella([[a['priorita'], a['cosaServe'], a.get('requisito', ''), a.get('esempioAccettabile', ''), a.get('destinatario', '')] for a in az],
                    intestazione=['Priorità', 'Cosa serve', 'Requisito', 'Evidenza accettabile', 'Destinatario'])
    else:
        doc.p('Nessuna azione aperta nel fascicolo.')
    return doc, k


# ---------------------------------------------------------------- avvio

def nome_file(tipo, k):
    stato = {'bozza': 'BOZZA', 'bozza_sospesa': 'BOZZA SOSPESA'}.get(k['stato'], '')
    # caratteri non ammessi da Windows; nomi lunghi accorciati per restare sotto il limite dei percorsi
    sicuro = lambda s, n: ''.join(ch for ch in s if ch not in '<>:"/\\|?*').strip()[:n].strip()
    return f"{tipo} {sicuro(k['cliente'], 40)} - {sicuro(k['bene'], 50)}{' - ' + stato if stato else ''} rev.{k['rev']} {k['data'].replace('/', '.')}.docx"


def main(argv=None):
    a = argparse.ArgumentParser(description=__doc__.split('\n')[0])
    a.add_argument('pratica')
    a.add_argument('--bene', action='append')
    a.add_argument('--perito')
    a.add_argument('--uscita')
    a.add_argument('--data', help='data dell’elaborato se il fascicolo non la indica (AAAA-MM-GG)')
    o = a.parse_args(argv)
    pr = Pratica(o.pratica)
    if pr.p['regime'] != 'CRI_4.0':
        sys.exit(f"regime {pr.p['regime']} non gestito")
    regime = leggi_json(Path(__file__).with_name('regime_CRI_4.0.json'))
    perito = leggi_json(o.perito) if o.perito else {}
    guida = guida_requisiti()
    oggi = data_it(o.data or dt.date.today().isoformat())
    uscita = Path(o.uscita) if o.uscita else pr.cartella / 'perizia'
    uscita.mkdir(parents=True, exist_ok=True)
    beni = o.bene or list(pr.schede)
    for codice in beni:
        if codice not in pr.schede:
            sys.exit(f'{codice}: bene non presente nel fascicolo')
        for fn, tipo in ((perizia, 'Perizia 4.0'), (analisi, 'Analisi Tecnica')):
            doc, k = fn(pr, codice, perito, regime, guida, oggi) if fn is analisi else fn(pr, codice, perito, regime, oggi)
            percorso = uscita / nome_file(tipo, k)
            if len(str(percorso.resolve())) > 255:
                sys.exit(f'percorso troppo lungo per Windows ({len(str(percorso.resolve()))} caratteri): {percorso}\n'
                         'usare --uscita con una cartella più corta o accorciare il nome della cartella della pratica')
            doc.salva(percorso)
            segnaposto = sum(par.text.count('[PLACEHOLDER') for par in doc.d.paragraphs) + sum(
                c.text.count('[PLACEHOLDER') for t in doc.d.tables for r in t.rows for c in r.cells)
            print(f'{percorso}  ({segnaposto} segnaposto)')


if __name__ == '__main__':
    main()
